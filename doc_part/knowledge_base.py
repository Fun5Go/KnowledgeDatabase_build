import os
import re
from typing import List, Dict, Optional
import json
import pdfplumber
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma


# =========================================================
# CONFIG
# =========================================================
FILE_PATH = r"C:\Users\FW\Desktop\FMEA_AI\Project_Phase\Codes\database\doc_part\FS6303220015R11.pdf"

PERSIST_DIR = r"./DATA/chroma_langchain_db"
COLLECTION_NAME = "fs_requirements"

# ---- PDF layout settings ----
HEADER_CROP = 90
FOOTER_CROP = 70
LEFT_ID_THRESHOLD = 160

# ---- start parsing from this heading ----
START_HEADING_PATTERN = re.compile(r"^\s*1\s+Introduction\s*$", re.IGNORECASE)

# ---- requirement ids ----
REQ_ID_CORE = r"(?:REQ|DRQ|LIM|CHO)_\d+(?:\*|\[[A-Za-z0-9_-]+\])?"
REQ_AT_LINE_START_PATTERN = re.compile(
    rf"^({REQ_ID_CORE})(?:\s+({REQ_ID_CORE}))?(?:\s+(.*))?$"
)
REQ_TOKEN_PATTERN = re.compile(REQ_ID_CORE)
# ---- normal text chunking ----
NORMAL_CHUNK_SIZE = 1200
NORMAL_CHUNK_OVERLAP = 150


# =========================================================
# UTILS
# =========================================================
def normalize_spaces(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_line(text: str) -> str:
    return normalize_spaces(text)


def is_noise_line(text: str) -> bool:
    low = text.lower().strip()

    if low in {
        "company confidential",
        "page",
        "date",
        "document title",
        "document id / release",
        "document id",
        "release",
    }:
        return True

    if re.fullmatch(r"\d+", text):
        return True

    if low.startswith("document title "):
        return True
    if low.startswith("document id "):
        return True

    return False


def clean_chunk_text(text: str) -> str:
    lines = [clean_line(line) for line in text.splitlines()]
    out = []

    for line in lines:
        if not line:
            continue
        if is_noise_line(line):
            continue
        out.append(line)

    return "\n".join(out).strip()


def split_text_with_overlap(text: str, chunk_size: int = 1200, overlap: int = 150) -> List[str]:
    """
    Simple character-based chunking for non-REQ text.
    Tries to split near paragraph/newline boundaries when possible.
    """
    text = text.strip()
    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    n = len(text)

    while start < n:
        end = min(start + chunk_size, n)

        if end < n:
            # try to cut at paragraph boundary
            para_break = text.rfind("\n\n", start, end)
            line_break = text.rfind("\n", start, end)
            sentence_break = text.rfind(". ", start, end)

            cut = max(para_break, line_break, sentence_break)
            if cut > start + chunk_size // 2:
                end = cut + (2 if cut == sentence_break else 0)

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= n:
            break

        start = max(end - overlap, start + 1)

    return chunks


# =========================================================
# PDF LOW-LEVEL HELPERS
# =========================================================
def group_words_to_lines(words: List[Dict], y_tolerance: float = 3) -> List[List[Dict]]:
    if not words:
        return []

    words = sorted(words, key=lambda w: (round(w["top"], 1), w["x0"]))
    lines = []
    current_line = [words[0]]
    current_top = words[0]["top"]

    for w in words[1:]:
        if abs(w["top"] - current_top) <= y_tolerance:
            current_line.append(w)
        else:
            lines.append(current_line)
            current_line = [w]
            current_top = w["top"]

    lines.append(current_line)
    return lines


def line_to_text_and_positions(line_words: List[Dict]):
    line_words = sorted(line_words, key=lambda x: x["x0"])
    text = " ".join(w["text"] for w in line_words).strip()
    x0 = min(w["x0"] for w in line_words)
    x1 = max(w["x1"] for w in line_words)
    top = min(w["top"] for w in line_words)
    return text, x0, x1, top


def extract_region_text(page, bbox, label: str = "") -> str:
    """
    Extract text from a page region for debug printing.
    bbox = (x0, top, x1, bottom)
    """
    try:
        region = page.crop(bbox)
        txt = region.extract_text() or ""
        txt = normalize_spaces(txt)
        return txt
    except Exception as e:
        return f"[ERROR extracting {label}: {e}]"


def print_pdf_header_footer_debug(file_path: str, max_pages: Optional[int] = None):
    """
    Print header/footer text of each page so you can verify crop settings.
    """
    print("\n" + "=" * 120)
    print("[DEBUG] HEADER / FOOTER CHECK")
    print("=" * 120)

    with pdfplumber.open(file_path) as pdf:
        total_pages = len(pdf.pages)
        pages_to_check = total_pages if max_pages is None else min(max_pages, total_pages)

        for page_idx in range(pages_to_check):
            page = pdf.pages[page_idx]

            header_bbox = (0, 0, page.width, HEADER_CROP)
            footer_bbox = (0, page.height - FOOTER_CROP, page.width, page.height)

            header_text = extract_region_text(page, header_bbox, "header")
            footer_text = extract_region_text(page, footer_bbox, "footer")

            print(f"\n--- PAGE {page_idx + 1} ---")
            print("[HEADER]")
            print(header_text if header_text else "[EMPTY]")
            print("[FOOTER]")
            print(footer_text if footer_text else "[EMPTY]")

def parse_requirement_line(text: str):
    """
    Parse lines like:
      REQ_5
      DRQ_202*
      REQ_169[PCR] Undervoltage protection shall be implemented on the iPS3
      DRQ_202* REQ_2
    Returns:
      {
        "ids": [...],
        "remainder": "..."
      }
    or None
    """
    m = REQ_AT_LINE_START_PATTERN.match(text.strip())
    if not m:
        return None

    id1 = m.group(1)
    id2 = m.group(2)
    remainder = (m.group(3) or "").strip()

    ids = [id1]
    if id2:
        ids.append(id2)

    return {
        "ids": ids,
        "remainder": remainder
    }


# =========================================================
# PDF PARSER
# =========================================================
def flush_normal_buffer(
    chunks: List[Document],
    normal_lines: List[str],
    file_path: str,
    pages: List[int],
):
    text = clean_chunk_text("\n".join(normal_lines))
    if not text:
        return

    split_parts = split_text_with_overlap(
        text,
        chunk_size=NORMAL_CHUNK_SIZE,
        overlap=NORMAL_CHUNK_OVERLAP,
    )

    for idx, part in enumerate(split_parts, start=1):
        chunks.append(
            Document(
                page_content=part,
                metadata={
                    "source": file_path,
                    "file_name": os.path.basename(file_path),
                    "pages": pages[:],
                    "type": "FS_section",
                    "chunk_mode": "normal",
                    "chunk_index": idx,
                },
            )
        )


def flush_req_buffer(
    chunks: List[Document],
    current_ids: Optional[List[str]],
    current_text_lines: List[str],
    file_path: str,
    pages: List[int],
):
    if not current_ids or not current_text_lines:
        return

    text = clean_chunk_text("\n".join(current_text_lines))
    if not text:
        return

    chunks.append(
        Document(
            page_content=text,
            metadata={
                "source": file_path,
                "file_name": os.path.basename(file_path),
                "pages": pages[:],
                "type": "FS",
                "chunk_mode": "requirement",
                "requirement_ids": current_ids[:],
                "primary_requirement_id": current_ids[0],
            },
        )
    )


def extract_fs_chunks_from_pdf(file_path: str) -> List[Document]:
    """
    Rules:
    1. Print/debug header/footer separately with helper
    2. Start parsing only after '1 Introduction'
    3. Before first REQ/DRQ/LIM/CHO -> normal chunking
    4. After REQ starts -> chunk by requirement block
    """
    chunks: List[Document] = []

    started = False
    in_requirement_mode = False

    normal_lines: List[str] = []
    normal_pages: List[int] = []

    current_ids: Optional[List[str]] = None
    current_req_lines: List[str] = []
    current_req_pages: List[int] = []

    with pdfplumber.open(file_path) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            page_no = page_idx + 1

            # crop out header/footer
            cropped = page.crop((0, HEADER_CROP, page.width, page.height - FOOTER_CROP))

            words = cropped.extract_words(
                use_text_flow=True,
                keep_blank_chars=False,
            )

            if not words:
                continue

            lines = group_words_to_lines(words, y_tolerance=3)

            for line_words in lines:
                raw_text, x0, _, _ = line_to_text_and_positions(line_words)
                text = clean_line(raw_text)

                if not text or is_noise_line(text):
                    continue

                # Wait until "1 Introduction"
                if not started:
                    if START_HEADING_PATTERN.match(text):
                        started = True
                        normal_lines.append(text)
                        if page_no not in normal_pages:
                            normal_pages.append(page_no)
                    continue

                # Check if this line is a requirement id line on the left
                # Check if this line starts a requirement block
                is_left_label = x0 < LEFT_ID_THRESHOLD
                parsed_req = parse_requirement_line(text)

                if is_left_label and parsed_req:
                    # switch normal -> requirement mode
                    if not in_requirement_mode:
                        flush_normal_buffer(chunks, normal_lines, file_path, normal_pages)
                        normal_lines = []
                        normal_pages = []
                        in_requirement_mode = True

                    # flush previous requirement
                    flush_req_buffer(chunks, current_ids, current_req_lines, file_path, current_req_pages)

                    current_ids = parsed_req["ids"]
                    current_req_lines = current_ids[:]
                    current_req_pages = [page_no]

                    # if the requirement line already contains body text, keep it
                    if parsed_req["remainder"]:
                        current_req_lines.append(parsed_req["remainder"])

                    continue

                # fallback requirement detection
                if is_left_label:
                    tokens = REQ_TOKEN_PATTERN.findall(text)
                    if tokens and len(tokens) <= 2 and text.replace(" ", "") in "".join(tokens).replace(" ", ""):
                        if not in_requirement_mode:
                            flush_normal_buffer(chunks, normal_lines, file_path, normal_pages)
                            normal_lines = []
                            normal_pages = []
                            in_requirement_mode = True

                        flush_req_buffer(chunks, current_ids, current_req_lines, file_path, current_req_pages)

                        current_ids = tokens
                        current_req_lines = tokens[:]
                        current_req_pages = [page_no]
                        continue

                # normal content handling
                if in_requirement_mode:
                    if current_ids:
                        current_req_lines.append(text)
                        if page_no not in current_req_pages:
                            current_req_pages.append(page_no)
                else:
                    normal_lines.append(text)
                    if page_no not in normal_pages:
                        normal_pages.append(page_no)

        # flush remaining buffers
        if in_requirement_mode:
            flush_req_buffer(chunks, current_ids, current_req_lines, file_path, current_req_pages)
        else:
            flush_normal_buffer(chunks, normal_lines, file_path, normal_pages)

    print(f"[INFO] Extracted {len(chunks)} chunks from PDF")
    return chunks


# =========================================================
# DOCX PARSER
# =========================================================
def extract_fs_chunks_from_docx(file_path: str) -> List[Document]:
    chunks: List[Document] = []

    started = False
    in_requirement_mode = False

    normal_lines: List[str] = []

    current_ids: Optional[List[str]] = None
    current_req_lines: List[str] = []

    doc = DocxDocument(file_path)

    for para in doc.paragraphs:
        text = clean_line(para.text)
        if not text or is_noise_line(text):
            continue

        if not started:
            if START_HEADING_PATTERN.match(text):
                started = True
                normal_lines.append(text)
            continue

        m = REQ_LINE_PATTERN.match(text)

        if m:
            if not in_requirement_mode:
                # flush normal chunks
                txt = clean_chunk_text("\n".join(normal_lines))
                if txt:
                    split_parts = split_text_with_overlap(
                        txt,
                        chunk_size=NORMAL_CHUNK_SIZE,
                        overlap=NORMAL_CHUNK_OVERLAP,
                    )
                    for idx, part in enumerate(split_parts, start=1):
                        chunks.append(
                            Document(
                                page_content=part,
                                metadata={
                                    "source": file_path,
                                    "file_name": os.path.basename(file_path),
                                    "type": "FS_section",
                                    "chunk_mode": "normal",
                                    "chunk_index": idx,
                                },
                            )
                        )
                normal_lines = []
                in_requirement_mode = True

            if current_ids and current_req_lines:
                req_text = clean_chunk_text("\n".join(current_req_lines))
                if req_text:
                    chunks.append(
                        Document(
                            page_content=req_text,
                            metadata={
                                "source": file_path,
                                "file_name": os.path.basename(file_path),
                                "type": "FS",
                                "chunk_mode": "requirement",
                                "requirement_ids": current_ids[:],
                                "primary_requirement_id": current_ids[0],
                            },
                        )
                    )

            current_ids = [g for g in m.groups() if g]
            current_req_lines = current_ids[:]
        else:
            if in_requirement_mode:
                if current_ids:
                    current_req_lines.append(text)
            else:
                normal_lines.append(text)

    # flush
    if in_requirement_mode:
        if current_ids and current_req_lines:
            req_text = clean_chunk_text("\n".join(current_req_lines))
            if req_text:
                chunks.append(
                    Document(
                        page_content=req_text,
                        metadata={
                            "source": file_path,
                            "file_name": os.path.basename(file_path),
                            "type": "FS",
                            "chunk_mode": "requirement",
                            "requirement_ids": current_ids[:],
                            "primary_requirement_id": current_ids[0],
                        },
                    )
                )
    else:
        txt = clean_chunk_text("\n".join(normal_lines))
        if txt:
            split_parts = split_text_with_overlap(
                txt,
                chunk_size=NORMAL_CHUNK_SIZE,
                overlap=NORMAL_CHUNK_OVERLAP,
            )
            for idx, part in enumerate(split_parts, start=1):
                chunks.append(
                    Document(
                        page_content=part,
                        metadata={
                            "source": file_path,
                            "file_name": os.path.basename(file_path),
                            "type": "FS_section",
                            "chunk_mode": "normal",
                            "chunk_index": idx,
                        },
                    )
                )

    print(f"[INFO] Extracted {len(chunks)} chunks from DOCX")
    return chunks


# =========================================================
# ENRICH FOR EMBEDDING
# =========================================================
def enrich_chunk_for_embedding(ch: Document) -> Document:
    source = ch.metadata.get("file_name", os.path.basename(ch.metadata.get("source", "")))
    doc_type = ch.metadata.get("type", "FS")
    chunk_mode = ch.metadata.get("chunk_mode", "")

    pages = ch.metadata.get("pages", [])
    if isinstance(pages, list) and pages:
        pages_str = ", ".join(map(str, pages))
    else:
        pages_str = ""

    req_ids = ", ".join(ch.metadata.get("requirement_ids", []))

    prefix_lines = [
        f"Document Type: {doc_type}",
        f"Source: {source}",
        f"Chunk Mode: {chunk_mode}",
    ]

    if pages_str:
        prefix_lines.append(f"Pages: {pages_str}")
    if req_ids:
        prefix_lines.append(f"Requirement IDs: {req_ids}")

    prefix_lines.append("Content:")

    content = "\n".join(prefix_lines) + "\n" + ch.page_content

    return Document(
        page_content=content.strip(),
        metadata=ch.metadata,
    )


def prepare_fs_chunks(file_path: str) -> List[Document]:
    lower = file_path.lower()

    if lower.endswith(".pdf"):
        chunks = extract_fs_chunks_from_pdf(file_path)
    elif lower.endswith(".docx"):
        chunks = extract_fs_chunks_from_docx(file_path)
    else:
        raise ValueError("Unsupported file type. Use PDF or DOCX.")

    enriched = [enrich_chunk_for_embedding(ch) for ch in chunks]
    print(f"[INFO] Total chunks ready for embedding: {len(enriched)}")
    return enriched


# =========================================================
# VECTOR STORE
# =========================================================
def create_embeddings(model_name: str = "sentence-transformers/all-MiniLM-L6-v2"):
    embeddings = HuggingFaceEmbeddings(
        model_name=model_name,
        model_kwargs={"device": "cpu"},
    )
    print("[INFO] Embeddings model initialized")
    return embeddings


def create_vector_store(
    embeddings,
    persist_dir: str = PERSIST_DIR,
    collection_name: str = COLLECTION_NAME,
):
    vector_store = Chroma(
        collection_name=collection_name,
        embedding_function=embeddings,
        persist_directory=persist_dir,
    )
    print("[INFO] Vector store initialized")
    return vector_store


def add_chunks_to_vector_store(vector_store, chunks: List[Document]):
    cleaned_chunks = []

    for ch in chunks:
        cleaned_metadata = {}
        for k, v in ch.metadata.items():
            if isinstance(v, (str, int, float, bool)) or v is None:
                cleaned_metadata[k] = v
            elif isinstance(v, list):
                cleaned_metadata[k] = json.dumps(v, ensure_ascii=False)
            else:
                cleaned_metadata[k] = str(v)

        cleaned_chunks.append(
            Document(
                page_content=ch.page_content,
                metadata=cleaned_metadata,
            )
        )

    ids = vector_store.add_documents(cleaned_chunks)
    print(f"[INFO] Added {len(ids)} chunks to vector store")
    return ids


# =========================================================
# MAIN
# =========================================================
def main():
    file_path = FILE_PATH

    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    # 1. debug header/footer
    print_pdf_header_footer_debug(file_path, max_pages=5)

    # 2. parse and chunk
    chunks = prepare_fs_chunks(file_path)

    # 3. preview chunks
    print("\n" + "=" * 120)
    print("[DEBUG] CHUNK PREVIEW")
    print("=" * 120)

    for i, ch in enumerate(chunks[:10], start=1):
        print(f"\n--- CHUNK {i} ---")
        print("METADATA:")
        print(ch.metadata)
        print("CONTENT:")
        print(ch.page_content[:1800])

    # 4. embedding + vector store
    embeddings = create_embeddings()
    vector_store = create_vector_store(embeddings)
    add_chunks_to_vector_store(vector_store, chunks)

    print("[DONE]")


if __name__ == "__main__":
    main()