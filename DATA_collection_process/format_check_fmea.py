import json
from pathlib import Path
from openpyxl import load_workbook

FMEA_DIR = Path(r"C:\Users\FW\Desktop\FMEA_AI\Project_Phase\DATA\RAW\FMEA")

BASE_DIR = Path(__file__).resolve().parent
JSON_PATH = BASE_DIR / "fmea_with_filename.json"

REQUIRED_HEADER = "FAILURE MODE EFFECT ANALYSIS"


VALID_SUFFIX = {".xlsx", ".xlsm"}
REQUIRED_WORDS = [
    "failure",
    "mode",
    "effect",
    "analysis"
]


def get_header_text(sheet):
    """
    Safely get first non-empty cell text from row 1
    Compatible with merged cells & read_only mode
    """
    for row in sheet.iter_rows(min_row=1, max_row=1):
        for cell in row:
            if cell.value:
                return str(cell.value)
    return None

def check_fmea_format(excel_path: Path) -> tuple[bool, str | None]:
    """
    return (is_valid, reason_if_invalid)
    """
    # 1. suffix check
    # if excel_path.suffix.lower() not in VALID_SUFFIX:
    #     return False, "unsupported file suffix"

    try:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
    except Exception:
        return False, "cannot open excel"

    # 2. sheet index = 1
    if len(wb.worksheets) <= 1:
        return False, "sheet1 not found"

    sheet = wb.worksheets[1]
    header = get_header_text(sheet)

    if not header:
        return False, "header empty"
    
    text = str(header).lower()

    if not all(word in text for word in REQUIRED_WORDS):
        return False, "header keywords missing"

    return True, None


def main():
    print(">>> FMEA format check started")

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        fmea_list = json.load(f)

    print(f">>> total files: {len(fmea_list)}")

    for r in fmea_list:
        filename = r.get("copiedFileName")
        r.setdefault("item", {})

        if not filename:
            r["item"]["format"] = False
            r["item"]["_format_reason"] = "missing filename"
            continue

        excel_path = FMEA_DIR / filename

        if not excel_path.exists():
            print(f"[MISS] {filename}")
            r["item"]["format"] = False
            r["item"]["_format_reason"] = "file not found"
            continue

        is_valid, reason = check_fmea_format(excel_path)
        r["item"]["format"] = is_valid

        if not is_valid:
            print(f"[BAD ] {filename} -> {reason}")
            r["item"]["_format_reason"] = reason
        else:
            print(f"[ OK ] {filename}")

    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(fmea_list, f, ensure_ascii=False, indent=2)

    print(">>> FMEA format check finished")


if __name__ == "__main__":
    main()