import os
import re
import json
from pathlib import Path
from typing import List, Dict, Optional, Tuple

import pdfplumber
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma



# =========================================================
# CONFIG
# =========================================================
DOC_PART_DIR = Path(__file__).resolve().parent
FILE_PATH = str(DOC_PART_DIR / "TS6303220029R09.pdf")

PERSIST_DIR = r"./DATA/chroma_langchain_db"
COLLECTION_NAME = "technical_specification"

# ---- PDF layout settings ----
HEADER_CROP = 90
LEFT_ID_THRESHOLD = 160
PREFIX_MERGE_MAX_GAP = 6.0

# ---- start parsing from this heading ----
START_HEADING_PATTERN = re.compile(r"^\s*1\s+Introduction\s*$", re.IGNORECASE)

# ---- requirement ids ----
REQ_ID_CORE = r"(?:REQ|DRQ|LIM|CHO|ASM)_\d+(?:\*|\[[A-Za-z0-9_-]+\])?"
REQ_AT_LINE_START_PATTERN = re.compile(
    rf"^({REQ_ID_CORE})(?:\s+({REQ_ID_CORE}))?(?:\s+(.*))?$"
)
REQ_TOKEN_PATTERN = re.compile(REQ_ID_CORE)
SECTION_HEADING_PATTERN = re.compile(
    r"^\s*\d+(?:\.\d+)+\s+[A-Z][^\n]*$"
)

# ---- normal text chunking ----
NORMAL_CHUNK_SIZE = 1200
NORMAL_CHUNK_OVERLAP = 150


# =========================================================
# UTILS
# =========================================================
def normalize_spaces(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_line(text: str) -> str:
    return normalize_spaces(text)


def is_noise_line(text: str) -> bool:
    low = text.lower().strip()

    if low in {
        "company confidential",
        "page",
        "date",
        "document title",
        "document id / release",
        "document id",
        "release",
    }:
        return True

    if re.fullmatch(r"\d+", low):
        return True

    if low.startswith("document title"):
        return True
    if low.startswith("document id"):
        return True
    if low.startswith("date "):
        return True

    return False


def merge_hyphenated_lines(lines: List[str]) -> List[str]:
    merged = []
    i = 0
    while i < len(lines):
        cur = lines[i].strip()
        if i < len(lines) - 1:
            nxt = lines[i + 1].strip()
            if cur.endswith("-") and nxt:
                merged.append(cur + nxt)
                i += 2
                continue
        merged.append(cur)
        i += 1
    return merged


def clean_chunk_text(text: str) -> str:
    lines = [clean_line(line) for line in text.splitlines()]
    lines = [line for line in lines if line and not is_noise_line(line)]
    lines = merge_hyphenated_lines(lines)
    return "\n".join(lines).strip()


def split_text_with_overlap(text: str, chunk_size: int = 1200, overlap: int = 150) -> List[str]:
    text = text.strip()
    if not text:
        return []

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    n = len(text)

    while start < n:
        end = min(start + chunk_size, n)

        if end < n:
            para_break = text.rfind("\n\n", start, end)
            line_break = text.rfind("\n", start, end)
            sentence_break = text.rfind(". ", start, end)
            cut = max(para_break, line_break, sentence_break)
            if cut > start + chunk_size // 2:
                end = cut + (2 if cut == sentence_break else 0)

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= n:
            break

        start = max(end - overlap, start + 1)

    return chunks


def parse_requirement_line(text: str):
    m = REQ_AT_LINE_START_PATTERN.match(text.strip())
    if not m:
        return None

    id1 = m.group(1)
    id2 = m.group(2)
    remainder = (m.group(3) or "").strip()

    ids = [id1]
    if id2:
        ids.append(id2)

    return {
        "ids": ids,
        "remainder": remainder
    }


def parse_requirement_ids(text: str) -> List[str]:
    return REQ_TOKEN_PATTERN.findall(text or "")


def is_requirement_label_text(text: str) -> bool:
    if not text:
        return False

    # 至少得有一个 requirement id
    tokens = parse_requirement_ids(text)
    if not tokens:
        return False

    # 允许的内容：
    # 1. requirement id，如 REQ_13 / DRQ_15 / CHO_35
    # 2. 引用标记，如 [11]
    # 3. 空白
    remainder = text

    # 去掉 requirement ids
    remainder = REQ_TOKEN_PATTERN.sub("", remainder)

    # 去掉引用标记 [11]
    remainder = re.sub(r"\[\d+\]", "", remainder)

    # 去掉所有空白
    remainder = re.sub(r"\s+", "", remainder)

    # 如果什么都不剩，说明整块左栏本质上就是 label 区
    return remainder == ""


def looks_like_rationale_label(text: str) -> bool:
    return text.strip().lower() == "rationale"


def looks_like_continuation(text: str) -> bool:
    if not text:
        return False
    t = text.strip()
    if not t:
        return False

    # 数字、标点、小写开头，通常像上一句的续行
    if re.match(r"^[0-9:;,\.\)\]-]", t):
        return True
    if re.match(r"^[a-z]", t):
        return True

    return False

def looks_like_section_heading(text: str) -> bool:
    if not text:
        return False

    t = clean_line(text)

    # avoid matching requirement ids accidentally
    if parse_requirement_ids(t):
        return False

    return bool(SECTION_HEADING_PATTERN.match(t))


def ends_with_hyphen(text_lines: List[str]) -> bool:
    for line in reversed(text_lines):
        s = line.strip()
        if s:
            return s.endswith("-")
    return False


def extract_requirement_only_text(chunk_text: str) -> str:
    lines = chunk_text.splitlines()
    req_parts = []

    for line in lines:
        line = clean_line(line)
        if not line:
            continue
        if line.startswith("Requirement:"):
            req_parts.append(line[len("Requirement:"):].strip())

    return clean_line(" ".join(req_parts))


def extract_rationale_only_text(chunk_text: str) -> str:
    lines = chunk_text.splitlines()
    rationale_parts = []

    for line in lines:
        line = clean_line(line)
        if not line:
            continue
        if line.startswith("Rationale:"):
            rationale_parts.append(line[len("Rationale:"):].strip())

    return clean_line(" ".join(rationale_parts))


# =========================================================
# PDF LOW-LEVEL HELPERS
# =========================================================
def group_words_to_lines(words: List[Dict], y_tolerance: float = 3.0) -> List[List[Dict]]:
    if not words:
        return []

    words = sorted(words, key=lambda w: (round(w["top"], 1), w["x0"]))
    lines = []
    current_line = [words[0]]
    current_top = words[0]["top"]

    for w in words[1:]:
        if abs(w["top"] - current_top) <= y_tolerance:
            current_line.append(w)
        else:
            lines.append(current_line)
            current_line = [w]
            current_top = w["top"]

    lines.append(current_line)
    return lines


def words_to_text(words: List[Dict]) -> str:
    if not words:
        return ""
    words = sorted(words, key=lambda x: x["x0"])
    return clean_line(" ".join(w["text"] for w in words))


def split_line_left_right(line_words: List[Dict], threshold: float = LEFT_ID_THRESHOLD) -> Tuple[str, str]:
    left_words = [w for w in line_words if w["x0"] < threshold]
    right_words = [w for w in line_words if w["x0"] >= threshold]

    left_text = words_to_text(left_words)
    right_text = words_to_text(right_words)
    return left_text, right_text


def line_to_info(line_words: List[Dict]) -> Dict:
    line_words = sorted(line_words, key=lambda x: x["x0"])
    text = " ".join(w["text"] for w in line_words).strip()
    x0 = min(w["x0"] for w in line_words)
    x1 = max(w["x1"] for w in line_words)
    top = min(w["top"] for w in line_words)
    bottom = max(w["bottom"] for w in line_words)

    return {
        "text": clean_line(text),
        "x0": x0,
        "x1": x1,
        "top": top,
        "bottom": bottom,
        "words": line_words,
    }


def lines_are_close(prev_line: Dict, curr_line: Dict, max_gap: float = PREFIX_MERGE_MAX_GAP) -> bool:
    gap = curr_line["top"] - prev_line["bottom"]
    return gap <= max_gap


def is_same_requirement_prefix_block(
    prev_line_info: Optional[Dict],
    curr_line_info: Dict,
    left_text: str,
    right_text: str,
    current_ids: Optional[List[str]],
    current_req_lines: List[str],
    max_gap: float = PREFIX_MERGE_MAX_GAP,
) -> bool:
    """
    判断当前这一行的 prefix 是否应该并入当前 requirement block，而不是新开 chunk。
    """
    if not current_ids:
        return False

    left_ids = parse_requirement_ids(left_text)
    if not left_ids or not is_requirement_label_text(left_text):
        return False

    close_to_prev = False
    if prev_line_info is not None:
        close_to_prev = lines_are_close(prev_line_info, curr_line_info, max_gap=max_gap)

    continuation_hint = (
        not right_text
        or looks_like_continuation(right_text)
        or ends_with_hyphen(current_req_lines)
    )

    return close_to_prev and continuation_hint


def extract_region_text(page, bbox, label: str = "") -> str:
    try:
        region = page.crop(bbox)
        txt = region.extract_text() or ""
        txt = normalize_spaces(txt)
        return txt
    except Exception as e:
        return f"[ERROR extracting {label}: {e}]"


def print_pdf_header_footer_debug(file_path: str, max_pages: Optional[int] = None):
    print("\n" + "=" * 120)
    print("[DEBUG] HEADER / FOOTER CHECK")
    print("=" * 120)

    with pdfplumber.open(file_path) as pdf:
        total_pages = len(pdf.pages)
        pages_to_check = total_pages if max_pages is None else min(max_pages, total_pages)

        for page_idx in range(pages_to_check):
            page = pdf.pages[page_idx]

            header_bbox = (0, 0, page.width, HEADER_CROP)
            footer_bbox = (0, page.height - 120, page.width, page.height)

            header_text = extract_region_text(page, header_bbox, "header")
            footer_text = extract_region_text(page, footer_bbox, "footer")

            print(f"\n--- PAGE {page_idx + 1} ---")
            print("[HEADER]")
            print(header_text if header_text else "[EMPTY]")
            print("[FOOTER]")
            print(footer_text if footer_text else "[EMPTY]")


# =========================================================
# FLUSH HELPERS
# =========================================================
def flush_normal_buffer(
    chunks: List[Document],
    normal_lines: List[str],
    file_path: str,
    pages: List[int],
):
    text = clean_chunk_text("\n".join(normal_lines))
    if not text:
        return

    split_parts = split_text_with_overlap(
        text,
        chunk_size=NORMAL_CHUNK_SIZE,
        overlap=NORMAL_CHUNK_OVERLAP,
    )

    for idx, part in enumerate(split_parts, start=1):
        chunks.append(
            Document(
                page_content=part,
                metadata={
                    "source": file_path,
                    "file_name": os.path.basename(file_path),
                    "pages": pages[:],
                    "type": "FS_section",
                    "chunk_mode": "normal",
                    "chunk_index": idx,
                },
            )
        )


def flush_req_buffer(
    chunks: List[Document],
    current_ids: Optional[List[str]],
    current_text_lines: List[str],
    file_path: str,
    pages: List[int],
):
    if not current_ids or not current_text_lines:
        return

    text = clean_chunk_text("\n".join(current_text_lines))
    if not text:
        return

    chunks.append(
        Document(
            page_content=text,
            metadata={
                "source": file_path,
                "file_name": os.path.basename(file_path),
                "pages": pages[:],
                "type": "FS",
                "chunk_mode": "requirement",
                "requirement_ids": current_ids[:],
                "primary_requirement_id": current_ids[0],
            },
        )
    )


def update_requirement_id_header(current_req_lines: List[str], current_ids: List[str]) -> List[str]:
    header = f"Requirement IDs: {', '.join(current_ids)}"
    if current_req_lines:
        current_req_lines[0] = header
    else:
        current_req_lines = [header]
    return current_req_lines


def append_requirement_text(current_req_lines: List[str], text: str):
    if not text:
        return

    if len(current_req_lines) >= 2 and current_req_lines[-1].startswith("Requirement: "):
        current_req_lines[-1] += " " + text
    else:
        current_req_lines.append(f"Requirement: {text}")


def append_rationale_text(current_req_lines: List[str], text: str):
    if not text:
        return

    if len(current_req_lines) >= 2 and current_req_lines[-1].startswith("Rationale: "):
        current_req_lines[-1] += " " + text
    else:
        current_req_lines.append(f"Rationale: {text}")


# =========================================================
# PDF PARSER
# =========================================================
def extract_fs_chunks_from_pdf(file_path: str) -> List[Document]:
    chunks: List[Document] = []

    started = False
    in_requirement_mode = False

    normal_lines: List[str] = []
    normal_pages: List[int] = []

    current_ids: Optional[List[str]] = None
    current_req_lines: List[str] = []
    current_req_pages: List[int] = []

    with pdfplumber.open(file_path) as pdf:
        for page_idx, page in enumerate(pdf.pages):
            page_no = page_idx + 1

            # 只裁 header，不裁 footer，避免把 rationale 裁掉
            cropped = page.crop((0, HEADER_CROP, page.width, page.height))

            words = cropped.extract_words(
                use_text_flow=False,
                keep_blank_chars=False,
                extra_attrs=["fontname", "size"],
            )

            if not words:
                continue

            lines = group_words_to_lines(words, y_tolerance=3.0)
            line_infos = [line_to_info(line_words) for line_words in lines]

            prev_line_info = None

            for line_info in line_infos:
                full_text = line_info["text"]
                left_text, right_text = split_line_left_right(line_info["words"], LEFT_ID_THRESHOLD)

                if not full_text or is_noise_line(full_text):
                    prev_line_info = line_info
                    continue

                if not started:
                    if START_HEADING_PATTERN.match(full_text):
                        started = True
                        normal_lines.append(full_text)
                        if page_no not in normal_pages:
                            normal_pages.append(page_no)
                    prev_line_info = line_info
                    continue

                left_ids = parse_requirement_ids(left_text)

                # -------------------------------------------------
                # Case 1: 左侧是 requirement prefix
                # -------------------------------------------------
                if left_ids and is_requirement_label_text(left_text):
                    # 如果像同一个 block 的附加 prefix，则并入当前 requirement
                    if is_same_requirement_prefix_block(
                        prev_line_info=prev_line_info,
                        curr_line_info=line_info,
                        left_text=left_text,
                        right_text=right_text,
                        current_ids=current_ids,
                        current_req_lines=current_req_lines,
                        max_gap=PREFIX_MERGE_MAX_GAP,
                    ):
                        for rid in left_ids:
                            if rid not in current_ids:
                                current_ids.append(rid)

                        current_req_lines = update_requirement_id_header(current_req_lines, current_ids)

                        if right_text:
                            append_requirement_text(current_req_lines, right_text)

                        if page_no not in current_req_pages:
                            current_req_pages.append(page_no)

                        prev_line_info = line_info
                        continue

                    # 否则是真正的新 requirement
                    if not in_requirement_mode:
                        flush_normal_buffer(chunks, normal_lines, file_path, normal_pages)
                        normal_lines = []
                        normal_pages = []
                        in_requirement_mode = True

                    flush_req_buffer(chunks, current_ids, current_req_lines, file_path, current_req_pages)

                    current_ids = left_ids[:]
                    current_req_pages = [page_no]
                    current_req_lines = [f"Requirement IDs: {', '.join(current_ids)}"]

                    if right_text:
                        append_requirement_text(current_req_lines, right_text)

                    prev_line_info = line_info
                    continue

                # -------------------------------------------------
                # Case 2: 左侧是 Rationale
                # -------------------------------------------------
                if looks_like_rationale_label(left_text):
                    if in_requirement_mode and current_ids:
                        if right_text:
                            append_rationale_text(current_req_lines, right_text)
                        else:
                            current_req_lines.append("Rationale:")

                        if page_no not in current_req_pages:
                            current_req_pages.append(page_no)
                    else:
                        normal_lines.append(full_text)
                        if page_no not in normal_pages:
                            normal_pages.append(page_no)

                    prev_line_info = line_info
                    continue

                 # -------------------------------------------------
                # Case 3: section heading -> end current requirement block
                # -------------------------------------------------
                if looks_like_section_heading(full_text):
                    if in_requirement_mode and current_ids:
                        flush_req_buffer(chunks, current_ids, current_req_lines, file_path, current_req_pages)

                        current_ids = None
                        current_req_lines = []
                        current_req_pages = []
                        in_requirement_mode = False

                    normal_lines.append(full_text)
                    if page_no not in normal_pages:
                        normal_pages.append(page_no)

                    prev_line_info = line_info
                    continue

                # -------------------------------------------------
                # Case 4: requirement 正文续行
                # -------------------------------------------------
                if in_requirement_mode and current_ids:
                    continuation = right_text if right_text else full_text

                    if continuation and not is_noise_line(continuation):
                        # 如果上一行是 Rationale，就继续拼到 Rationale
                        if current_req_lines and current_req_lines[-1].startswith("Rationale:"):
                            append_rationale_text(current_req_lines, continuation)
                        else:
                            append_requirement_text(current_req_lines, continuation)

                        if page_no not in current_req_pages:
                            current_req_pages.append(page_no)

                    prev_line_info = line_info
                    continue


        # flush remaining buffers
        if in_requirement_mode:
            flush_req_buffer(chunks, current_ids, current_req_lines, file_path, current_req_pages)
        else:
            flush_normal_buffer(chunks, normal_lines, file_path, normal_pages)

    print(f"[INFO] Extracted {len(chunks)} chunks from PDF")
    return chunks


# =========================================================
# DOCX PARSER
# =========================================================
def extract_fs_chunks_from_docx(file_path: str) -> List[Document]:
    chunks: List[Document] = []

    started = False
    in_requirement_mode = False

    normal_lines: List[str] = []

    current_ids: Optional[List[str]] = None
    current_req_lines: List[str] = []

    doc = DocxDocument(file_path)

    for para in doc.paragraphs:
        text = clean_line(para.text)
        if not text or is_noise_line(text):
            continue

        if not started:
            if START_HEADING_PATTERN.match(text):
                started = True
                normal_lines.append(text)
            continue

        parsed = parse_requirement_line(text)

        if parsed:
            if not in_requirement_mode:
                txt = clean_chunk_text("\n".join(normal_lines))
                if txt:
                    split_parts = split_text_with_overlap(
                        txt,
                        chunk_size=NORMAL_CHUNK_SIZE,
                        overlap=NORMAL_CHUNK_OVERLAP,
                    )
                    for idx, part in enumerate(split_parts, start=1):
                        chunks.append(
                            Document(
                                page_content=part,
                                metadata={
                                    "source": file_path,
                                    "file_name": os.path.basename(file_path),
                                    "type": "FS_section",
                                    "chunk_mode": "normal",
                                    "chunk_index": idx,
                                },
                            )
                        )
                normal_lines = []
                in_requirement_mode = True

            flush_req_buffer(chunks, current_ids, current_req_lines, file_path, [])

            current_ids = parsed["ids"][:]
            current_req_lines = [f"Requirement IDs: {', '.join(current_ids)}"]

            if parsed["remainder"]:
                append_requirement_text(current_req_lines, parsed["remainder"])

        else:
            if in_requirement_mode and current_ids:
                if looks_like_rationale_label(text):
                    current_req_lines.append("Rationale:")
                else:
                    if current_req_lines and current_req_lines[-1].startswith("Rationale:"):
                        append_rationale_text(current_req_lines, text)
                    else:
                        append_requirement_text(current_req_lines, text)
            else:
                normal_lines.append(text)

    if in_requirement_mode:
        flush_req_buffer(chunks, current_ids, current_req_lines, file_path, [])
    else:
        txt = clean_chunk_text("\n".join(normal_lines))
        if txt:
            split_parts = split_text_with_overlap(
                txt,
                chunk_size=NORMAL_CHUNK_SIZE,
                overlap=NORMAL_CHUNK_OVERLAP,
            )
            for idx, part in enumerate(split_parts, start=1):
                chunks.append(
                    Document(
                        page_content=part,
                        metadata={
                            "source": file_path,
                            "file_name": os.path.basename(file_path),
                            "type": "FS_section",
                            "chunk_mode": "normal",
                            "chunk_index": idx,
                        },
                    )
                )

    print(f"[INFO] Extracted {len(chunks)} chunks from DOCX")
    return chunks


# =========================================================
# ENRICH FOR EMBEDDING
# =========================================================
def enrich_chunk_for_embedding(ch: Document) -> Document:
    source = ch.metadata.get("file_name", os.path.basename(ch.metadata.get("source", "")))
    doc_type = ch.metadata.get("type", "FS")
    chunk_mode = ch.metadata.get("chunk_mode", "")

    pages = ch.metadata.get("pages", [])
    if isinstance(pages, list) and pages:
        pages_str = ", ".join(map(str, pages))
    else:
        pages_str = ""

    req_ids = ", ".join(ch.metadata.get("requirement_ids", []))

    prefix_lines = [
        f"Document Type: {doc_type}",
        f"Source: {source}",
        f"Chunk Mode: {chunk_mode}",
    ]

    if pages_str:
        prefix_lines.append(f"Pages: {pages_str}")
    if req_ids:
        prefix_lines.append(f"Requirement IDs: {req_ids}")

    prefix_lines.append("Content:")

    content = "\n".join(prefix_lines) + "\n" + ch.page_content

    return Document(
        page_content=content.strip(),
        metadata=ch.metadata,
    )


def prepare_fs_chunks(file_path: str) -> List[Document]:
    lower = file_path.lower()

    if lower.endswith(".pdf"):
        chunks = extract_fs_chunks_from_pdf(file_path)
    elif lower.endswith(".docx"):
        chunks = extract_fs_chunks_from_docx(file_path)
    else:
        raise ValueError("Unsupported file type. Use PDF or DOCX.")

    enriched = [enrich_chunk_for_embedding(ch) for ch in chunks]
    print(f"[INFO] Total chunks ready for embedding: {len(enriched)}")
    return enriched


# =========================================================
# VECTOR STORE
# =========================================================
def create_embeddings(model_name: str = "sentence-transformers/all-MiniLM-L6-v2"):
    embeddings = HuggingFaceEmbeddings(
        model_name=model_name,
        model_kwargs={"device": "cpu"},
    )
    print("[INFO] Embeddings model initialized")
    return embeddings


def create_vector_store(
    embeddings,
    persist_dir: str = PERSIST_DIR,
    collection_name: str = COLLECTION_NAME,
):
    vector_store = Chroma(
        collection_name=collection_name,
        embedding_function=embeddings,
        persist_directory=persist_dir,
    )
    print("[INFO] Vector store initialized")
    return vector_store


def add_chunks_to_vector_store(vector_store, chunks: List[Document]):
    cleaned_chunks = []

    for ch in chunks:
        cleaned_metadata = {}
        for k, v in ch.metadata.items():
            if isinstance(v, (str, int, float, bool)) or v is None:
                cleaned_metadata[k] = v
            elif isinstance(v, list):
                cleaned_metadata[k] = json.dumps(v, ensure_ascii=False)
            else:
                cleaned_metadata[k] = str(v)

        cleaned_chunks.append(
            Document(
                page_content=ch.page_content,
                metadata=cleaned_metadata,
            )
        )

    ids = vector_store.add_documents(cleaned_chunks)
    print(f"[INFO] Added {len(ids)} chunks to vector store")
    return ids


# =========================================================
# MAIN
# =========================================================
def main():
    file_path = FILE_PATH

    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    # 1. debug header/footer
    print_pdf_header_footer_debug(file_path, max_pages=5)

    # 2. parse and chunk
    chunks = prepare_fs_chunks(file_path)

    # 3. preview chunks
    print("\n" + "=" * 120)
    print("[DEBUG] CHUNK PREVIEW")
    print("=" * 120)

    for i, ch in enumerate(chunks[:20], start=1):
        print(f"\n--- CHUNK {i} ---")
        print("METADATA:")
        print(ch.metadata)
        print("CONTENT:")
        print(ch.page_content[:2000])

    # 4. embedding + vector store
    embeddings = create_embeddings()
    vector_store = create_vector_store(embeddings)
    add_chunks_to_vector_store(vector_store, chunks)

    print("[DONE]")


if __name__ == "__main__":
    main()
